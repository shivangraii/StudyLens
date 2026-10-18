from __future__ import annotations


def parse_pdf(path: str) -> list[tuple[int, str]]:
    """Return [(page_number, text)] using PyMuPDF."""
    import pymupdf

    pages = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append((i, text))
    if not pages:
        raise ValueError("No extractable text in this PDF (scanned?). Run OCR first.")
    return pages


def parse_pptx(path: str) -> list[tuple[int, str]]:
    """Return [(slide_number, text)] including speaker notes."""
    from pptx import Presentation

    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notes: " + notes)
        if parts:
            slides.append((i, "\n".join(parts)))
    if not slides:
        raise ValueError("No text found in this presentation.")
    return slides
